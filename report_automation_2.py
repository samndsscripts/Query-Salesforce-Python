import time
import csv
from simple_salesforce import Salesforce
import json
from tabulate import tabulate
from openpyxl import load_workbook
import pandas as pd
import xlwings as xw
import re
from collections import Counter
from rapidfuzz import process

# ----- Salesforce login -----
sf = Salesforce(
    username='samuelcooper@ndspro.com',
    password='Summer@NDS2025',
    security_token='zjU2IJAfQmx6zDxgOj3aLkyPQ',
    instance_url='https://nds.my.salesforce.com'
)

# --- Fetch Salesforce report ---
report_id = "00OUI00000EsGR72AN"

try:
    report_data = sf.restful(f'analytics/reports/{report_id}', params={'includeDetails': 'true'})
    print("✅ Successfully fetched report data!")
    print("Report name:", report_data.get('reportMetadata', {}).get('name', 'Unknown'))
    print("Number of rows:", len(report_data.get('factMap', {}).get('T!T', {}).get('rows', [])))
except Exception as e:
    print("❌ Error fetching report:", e)

# Limit for testing — process first 5 rows
report_rows = report_data.get('factMap', {}).get('T!T', {}).get('rows', [])[:5]

# --- Load supplier table from Excel ---
wb = xw.Book(r"C:\Users\emp35107\OneDrive - NORMA Group\Documents\PPMs.xlsx")
sheet_sup = wb.sheets['Class Site Supplier']
table_range = sheet_sup.tables['Table3'].range
supplier_df = pd.DataFrame(table_range.value[1:], columns=table_range.value[0])

# --- Clean and align lists ---
supplier_df = supplier_df.dropna(subset=['Stock Code'])  # keep only rows with valid Stock Codes
stock_codes_list = supplier_df['Stock Code'].astype(str).tolist()
man_sites_list   = supplier_df['Mfg Plant Name'].fillna("").astype(str).tolist()
suppliers_list   = supplier_df['Supplier'].fillna("").astype(str).tolist()

# --- Normalize helper ---
def normalize(text):
    text = str(text).lower()
    text = re.sub(r"[^a-z0-9\s]", " ", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()

# --- Get reference to output table ---
sheet_out = wb.sheets['test']
table_out = sheet_out.tables['test']
output_range = table_out.range
output_df = pd.DataFrame(output_range.value[1:], columns=output_range.value[0])

# --- Process each Salesforce row ---
for row in report_rows:
    cells = row.get('dataCells', [])
    description = str(cells[4].get('label', ''))

    # --- Match to supplier/manufacturer ---
    normalized_description = normalize(description)
    normalized_stock_codes = [normalize(code) for code in stock_codes_list]
    matches = process.extract(normalized_description, normalized_stock_codes, limit=1)

    most_common_source = "N/A"
    if matches:
        best_match = matches[0][0]
        try:
            match_index = normalized_stock_codes.index(best_match)
            if match_index < len(suppliers_list):
                supplier = suppliers_list[match_index].strip()
                most_common_source = supplier if supplier else man_sites_list[match_index]
            else:
                print(f"⚠️ Warning: Match index {match_index} out of range for lists.")
        except Exception as e:
            print(f"⚠️ Match error for description '{description}': {e}")

    # --- Create new row ---
    new_row = [
        cells[0].get('label', ''),  # Opened Date
        cells[1].get('label', ''),  # Case Reason
        cells[2].get('label', ''),  # Case Owner
        cells[3].get('label', ''),  # Case Number
        description,                # Description
        '',                         # Quantity
        cells[5].get('label', ''),  # RMA Value
        cells[6].get('label', ''),  # Case Category
        cells[7].get('label', ''),  # Account Name
        '',                         # Comments
        cells[8].get('label', ''),  # Contact Type
        cells[9].get('label', ''),  # Shipping Whse
        most_common_source          # Source
    ]
  # --- Append new row into table (extend table range) ---
    table_range = table_out.range
    next_row = table_range.last_cell.row + 1

    # Write new row values
    sheet_out.range(f"A{next_row}").value = new_row

    # Resize table to include new row
    new_last_cell = sheet_out.range(f"A{next_row}").end('right')
    new_table_range = sheet_out.range(table_range.address.split(':')[0] + ':' + new_last_cell.address)
    table_out.resize(new_table_range)w

    print(f"✅ Added row for Case Number {cells[3].get('label', '')} with Source '{most_common_source}'")

print("\n✅ Finished processing first 5 rows.")
