import time
import csv
from simple_salesforce import Salesforce
import json
from tabulate import tabulate
from openpyxl import load_workbook
import pandas as pd
import xlwings as xw
import re
from collections import Counter
from rapidfuzz import process

# ----- Salesforce login -----
sf = Salesforce(
    username='samuelcooper@ndspro.com',
    password='Summer@NDS2025',
    security_token='zjU2IJAfQmx6zDxgOj3aLkyPQ',
    instance_url='https://nds.my.salesforce.com'
)

# --- Fetch Salesforce report ---
report_id = "00OUI00000EsGR72AN"
try:
    report_data = sf.restful(f'analytics/reports/{report_id}', params={'includeDetails': 'true'})
except Exception as e:
    print("❌ Error fetching report:", e)

# --- Limit for testing — first 5 rows ---
report_rows = report_data.get('factMap', {}).get('T!T', {}).get('rows', [])[:5]

# --- Load supplier table from Excel ---
wb = xw.Book(r"C:\Users\emp35107\OneDrive - NORMA Group\Documents\PPMs.xlsx")
sheet_sup = wb.sheets['Class Site Supplier']
table_range = sheet_sup.tables['Table3'].range
supplier_df = pd.DataFrame(table_range.value[1:], columns=table_range.value[0])

# --- Clean and align lists ---
supplier_df = supplier_df.dropna(subset=['Stock Code'])
stock_codes_list = supplier_df['Stock Code'].astype(str).tolist()
man_sites_list   = supplier_df['Mfg Plant Name'].fillna("").astype(str).tolist()
suppliers_list   = supplier_df['Supplier'].fillna("").astype(str).tolist()

# --- Normalize helper ---
def normalize(text):
    text = str(text).lower()
    text = re.sub(r"[^a-z0-9\s]", " ", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()

# --- Get reference to output table ---
sheet_out = wb.sheets['test']
table_out = sheet_out.tables['test']
output_range = table_out.range
output_df = pd.DataFrame(output_range.value[1:], columns=output_range.value[0])

# Track existing case numbers as integers
existing_cases = set(
    output_df['Case Number'].dropna().apply(lambda x: int(float(x))).tolist()
)

# --- Process each Salesforce row ---
new_rows_to_add = []
rows_processed = 0

for row in report_rows:
    rows_processed += 1
    cells = row.get('dataCells', [])

    # Convert Salesforce Case Number to integer
    try:
        case_number = int(float(cells[3].get('label', 0)))
    except:
        print(f"⚠️ Skipping row with invalid Case Number: {cells[3].get('label', '')}")
        continue

    # Skip if already in table
    if case_number in existing_cases:
        print(f"Row not added — Case Number {case_number} already in table")
        continue

    description = str(cells[4].get('label', ''))

    # --- Match to supplier/manufacturer ---
    normalized_description = normalize(description)
    normalized_stock_codes = [normalize(code) for code in stock_codes_list]
    matches = process.extract(normalized_description, normalized_stock_codes, limit=1)

    most_common_source = "N/A"
    if matches:
        best_match = matches[0][0]
        try:
            match_index = normalized_stock_codes.index(best_match)
            supplier = suppliers_list[match_index].strip() if match_index < len(suppliers_list) else ""
            most_common_source = supplier if sup
